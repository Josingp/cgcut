"""pytest: 타임코드 계산 / 리스트 파싱 / 실제 컷(ffmpeg 필요, 없으면 건너뜀)."""
from __future__ import annotations

import shutil
import subprocess
import sys
from fractions import Fraction
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import cgcut  # noqa: E402
from cgcut import TC  # noqa: E402

HAVE_FFMPEG = shutil.which("ffmpeg") and shutil.which("ffprobe")


# --------------------------------------------------------------------------- TC
def test_tc_parse_variants():
    assert TC.parse("01:04:24:08") == TC(1, 4, 24, 8)
    assert TC.parse("01:04:24;08") == TC(1, 4, 24, 8, drop=True)
    assert TC.parse("01:04:24.08") == TC(1, 4, 24, 8)
    assert TC.parse("01:04:24") == TC(1, 4, 24, 0)
    assert TC.parse("TC 01:04:24:08 부터") == TC(1, 4, 24, 8)
    assert TC.parse("없음") is None
    rng = TC.find_all("01:36:10:07\n~\n01:36:20:03")
    assert rng == [TC(1, 36, 10, 7), TC(1, 36, 20, 3)]


def test_tc_frames_roundtrip_ndf():
    for base in (24, 25, 30, 60):
        for n in (0, 1, 23, 24, 1439, 86399, 86400 * base // 24 + 7):
            assert TC.from_frames(n, base).to_frames(base) == n
    assert TC(1, 0, 0, 0).to_frames(24) == 86400
    assert str(TC.from_frames(86400 + 48, 24)) == "01:00:02:00"


def test_tc_frames_roundtrip_df():
    # 29.97 드롭프레임: 00:00:59;29 = 1799, 다음 프레임은 00:01:00;02 (=1800)
    assert TC(0, 0, 59, 29, drop=True).to_frames(30) == 1799
    assert TC(0, 1, 0, 2, drop=True).to_frames(30) == 1800
    assert str(TC.from_frames(1800, 30, drop=True)) == "00:01:00;02"
    assert TC(0, 10, 0, 0, drop=True).to_frames(30) == 17982
    for n in (0, 1797, 1798, 17981, 17982, 107892, 2589407):
        assert TC.from_frames(n, 30, drop=True).to_frames(30) == n


def test_tc_base():
    assert cgcut.tc_base(Fraction(24000, 1001)) == 24
    assert cgcut.tc_base(Fraction(30000, 1001)) == 30
    assert cgcut.tc_base(Fraction(25)) == 25


# --------------------------------------------------------------------------- 이름 추출
def test_ep_from_name():
    f = cgcut.ep_from_name
    assert f("EP04") == 4
    assert f("쩔해주세요_EP 12_fine.mov") == 12
    assert f("힘내라여대리_3화_편집본.mp4") == 3
    assert f("S01E07.mkv") == 7
    assert f("제12회_본방.mov") == 12
    assert f("master_260916.mov") is None


def test_clip_tokens():
    assert cgcut.clip_tokens("A027 C023") == ["A027C023"]
    assert cgcut.clip_tokens("A028 C022,\nA028 C023") == ["A028C022", "A028C023"]
    assert cgcut.clip_tokens("A019 C036\nor\nB019 C032") == ["A019C036", "B019C032"]
    assert cgcut.clip_tokens("A008 C013\n~A008 C017\n중 활용컷") == [f"A008C{c:03d}" for c in range(13, 18)]
    assert cgcut.clip_tokens("C015 C0006, C015 C0007") == ["C015C0006", "C015C0007"]
    assert cgcut.clip_tokens("휴대폰 든 직원 손") == []


def test_find_clip_files(tmp_path: Path):
    for name in ["A027C023_250913_R1AB.mov", "A027_C023_0913.R3D", "A027_09131234_C023.braw",
                 "A027C0230_x.mov", "B027C023.mxf", "A0027C023.mov", "A027C023.mov.sidecar"]:
        (tmp_path / name).write_bytes(b"x")
    found = sorted(p.name for p in cgcut.find_clip_files(tmp_path, "A027C023"))
    assert found == ["A0027C023.mov", "A027C023.mov.sidecar", "A027C023_250913_R1AB.mov",
                     "A027_09131234_C023.braw", "A027_C023_0913.R3D"]


# --------------------------------------------------------------------------- xlsx 파싱
@pytest.fixture(scope="module")
def sample_dir(tmp_path_factory) -> Path:
    import make_sample
    d = tmp_path_factory.mktemp("sample")
    make_sample.make_xlsx(d / "CG_list_sample.xlsx")
    if HAVE_FFMPEG:
        make_sample.make_video(d / "EP01_sample.mov", "01:00:00:00")
        make_sample.make_video(d / "EP02_sample.mp4", None)
    return d


def test_parse_workbook(sample_dir: Path):
    rows = cgcut.parse_workbook(sample_dir / "CG_list_sample.xlsx")
    assert len(rows) == 11
    by = {(r.sheet, r.no): r for r in rows}
    assert by[("EP01", "1")].tc_in == TC(1, 0, 2, 0) and by[("EP01", "1")].ep == 1
    assert by[("EP01", "3")].tc_out == TC(1, 0, 12, 5)
    assert by[("EP01", "7")].tc_in is None
    assert by[("EP02", "2-1")].label == "EP02_2-1"
    assert by[("EP01", "2")].desc.startswith("샷4")


def test_parse_headerless_sheet(tmp_path: Path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "memo"
    ws["B3"] = "뭔가"
    ws["C3"] = "01:00:01:00"
    ws["C4"] = "01:00:02:12"
    ws["D4"] = "비고 12:30"
    wb.save(tmp_path / "x.xlsx")
    rows = cgcut.parse_workbook(tmp_path / "x.xlsx")
    assert [r.tc_in for r in rows] == [TC(1, 0, 1, 0), TC(1, 0, 2, 12)]


# --------------------------------------------------------------------------- 실제 컷
def _gray_frames(path: Path, w: int, h: int) -> list[float]:
    raw = subprocess.run(["ffmpeg", "-v", "error", "-i", str(path), "-an", "-f", "rawvideo", "-pix_fmt", "gray", "-"],
                         capture_output=True, check=True).stdout
    n = len(raw) // (w * h)
    return [sum(raw[i * w * h:(i + 1) * w * h]) / (w * h) for i in range(n)]


def _stream_tc(path: Path) -> str:
    out = subprocess.run(["ffprobe", "-v", "error", "-select_streams", "v:0", "-show_entries", "stream_tags=timecode",
                          "-of", "csv=p=0", str(path)], capture_output=True, text=True).stdout
    return out.strip()


@pytest.mark.skipif(not HAVE_FFMPEG, reason="ffmpeg 없음")
def test_cut_shot_mode(sample_dir: Path, tmp_path: Path):
    import make_sample as ms
    out = tmp_path / "out"
    rc = cgcut.main([str(sample_dir / "CG_list_sample.xlsx"), "--video-dir", str(sample_dir), "-o", str(out), "--codec", "prores"])
    assert rc == 0
    b = [0] + ms.boundaries()
    # 라벨 → (시작 프레임, 프레임 수)
    expected = {"EP01_1": (48, 72), "EP01_2": (150, 96), "EP01_3": (246, 48), "EP01_4": (120, 30),
                "EP01_8": (0, 48), "EP01_9": (666, 56), "EP02_1": (48, 72), "EP02_2-1": (426, 40)}
    files = sorted(out.glob("*.mov"))
    assert {"_".join(f.name.split("_")[:2]) for f in files} == set(expected)
    for f in files:
        key = "_".join(f.name.split("_")[:2])
        start, n = expected[key]
        means = _gray_frames(f, ms.W, ms.H)
        seg = next(i for i in range(len(b) - 1) if b[i] <= start < b[i + 1])
        assert len(means) == n, f"{key}: 프레임 수 {len(means)} != {n}"
        assert max(means) - min(means) < 3, f"{key}: 샷 경계를 넘음"
        assert abs(means[0] - ms.GRAYS[seg]) < 6, f"{key}: 엉뚱한 샷"
        assert _stream_tc(f) == str(TC.from_frames(86400 + start, 24)), f"{key}: 타임코드 메타데이터"
    report = (out / "cgcut_report.csv").read_text(encoding="utf-8-sig").splitlines()
    assert len(report) == 12
    assert sum("건너뜀" in line for line in report) == 3


@pytest.mark.skipif(not HAVE_FFMPEG, reason="ffmpeg 없음")
def test_cut_fixed_handles_png_copy(sample_dir: Path, tmp_path: Path):
    import make_sample as ms
    xlsx, vid = str(sample_dir / "CG_list_sample.xlsx"), str(sample_dir / "EP01_sample.mov")
    # fixed: no.2(200) 앞 1초 뒤 2초 → 176~247 (72f)
    out = tmp_path / "fixed"
    assert cgcut.main([xlsx, "--sheet", "EP01", "--video", vid, "-o", str(out), "--mode", "fixed", "--before", "1", "--after", "2", "--codec", "h264"]) == 0
    f = next(out.glob("EP01_2_*.mp4"))
    assert f.name == "EP01_2_01-00-07-08.mp4" and len(_gray_frames(f, ms.W, ms.H)) == 72
    # handles 12: no.1(48~119) → 36~131 (96f)
    out = tmp_path / "handles"
    assert cgcut.main([xlsx, "--sheet", "EP01", "--video", vid, "-o", str(out), "--handles", "12", "--codec", "h264"]) == 0
    f = next(out.glob("EP01_1_*.mp4"))
    assert f.name == "EP01_1_01-00-01-12.mp4" and len(_gray_frames(f, ms.W, ms.H)) == 96
    # png 시퀀스
    out = tmp_path / "png"
    assert cgcut.main([xlsx, "--sheet", "EP01", "--video", vid, "-o", str(out), "--codec", "png"]) == 0
    pngs = sorted((out / "EP01_4_01-00-05-00").glob("*.png"))
    assert len(pngs) == 30 and pngs[0].name.endswith(".001001.png")
    # copy(키프레임 단위): 실행만 확인
    out = tmp_path / "copy"
    assert cgcut.main([xlsx, "--sheet", "EP01", "--video", vid, "-o", str(out), "--codec", "copy"]) == 0
    assert (out / "EP01_1_01-00-02-00.mov").exists()
    # dry-run은 파일을 만들지 않음
    out = tmp_path / "dry"
    assert cgcut.main([xlsx, "--video", vid, "-o", str(out), "--dry-run"]) == 0
    assert list(out.glob("*.mov")) == [] and (out / "cgcut_report.csv").exists()


@pytest.mark.skipif(not HAVE_FFMPEG, reason="ffmpeg 없음")
def test_start_tc_override(sample_dir: Path, tmp_path: Path):
    """--start-tc 00:59:58:00 이면 리스트 TC 01:00:02:00 은 파일 4초(96프레임) 지점."""
    import make_sample as ms
    xlsx, vid = str(sample_dir / "CG_list_sample.xlsx"), str(sample_dir / "EP01_sample.mov")
    out = tmp_path / "stc"
    assert cgcut.main([xlsx, "--sheet", "EP01", "--video", vid, "-o", str(out), "--start-tc", "00:59:58:00", "--codec", "h264", "--mode", "fixed", "--before", "0", "--after", "1"]) == 0
    f = next(out.glob("EP01_1_*.mp4"))
    # 01:00:02:00 - 00:59:58:00 = 4초 = 96프레임 → 둘째 샷(48~119) 구간. 파일명 TC는 start-tc 기준.
    assert f.name == "EP01_1_01-00-02-00.mp4"
    means = _gray_frames(f, ms.W, ms.H)
    assert len(means) == 24 and abs(means[0] - ms.GRAYS[1]) < 6


def test_clips_mode(sample_dir: Path, tmp_path: Path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "List"
    ws.append(["No.", "회차", "EP", "S#", "C#", "LENS", "Scan Name", "Image", "Source", "CG description", "NOTE"])
    ws.append([1, "8회차", 1, "1-2", 6, "89mm", "A027 C023", None, None, "빈백 합성", None])
    ws.append([2, "7회차", 3, "7", 3, None, "A026 C025\nor\nB026 C025", None, None, "REMOVE", None])
    ws.append([3, "5회차", 50, "110", None, None, None, None, None, "자막", None])
    xlsx = tmp_path / "list.xlsx"
    wb.save(xlsx)
    cam = tmp_path / "cam"
    cam.mkdir()
    (cam / "A027C023_250913_R1AB.mov").write_bytes(b"a" * 10)
    (cam / "B026C025_250912.mov").write_bytes(b"b" * 10)
    out = tmp_path / "out"
    assert cgcut.main([str(xlsx), "--clips-dir", str(cam), "-o", str(out)]) == 0
    got = sorted(p.name for p in (out / "clips").iterdir())
    assert got == ["EP01_1_A027C023_250913_R1AB.mov", "EP03_2_B026C025_250912.mov"]
    rep = (out / "cgcut_clips_report.csv").read_text(encoding="utf-8-sig")
    assert "A026C025" in rep and "못 찾음" in rep


@pytest.mark.skipif(not HAVE_FFMPEG, reason="ffmpeg 없음")
def test_windowed_scene_scan(sample_dir: Path, tmp_path: Path):
    """샷이 적으면 전체 대신 TC 주변만 검사하는 경로."""
    import make_sample as ms
    out = tmp_path / "win"
    assert cgcut.main([str(sample_dir / "CG_list_sample.xlsx"), "--sheet", "EP02", "--video", str(sample_dir / "EP02_sample.mp4"),
                       "-o", str(out), "--max-before", "3", "--max-after", "3", "--codec", "h264"]) == 0
    assert not (out / ".cache").exists()  # 전체 스캔 안 함
    f = next(out.glob("EP02_1_*.mp4"))
    assert f.name == "EP02_1_01-00-02-00.mp4" and len(_gray_frames(f, ms.W, ms.H)) == 72


@pytest.mark.skipif(not HAVE_FFMPEG, reason="ffmpeg 없음")
def test_drop_mode(sample_dir: Path, tmp_path: Path):
    import shutil as sh
    d = tmp_path / "drop"
    d.mkdir()
    sh.copy(sample_dir / "CG_list_sample.xlsx", d / "list.xlsx")
    assert cgcut.main(["--drop", str(d / "list.xlsx"), str(sample_dir / "EP01_sample.mov"), "--codec", "h264", "--sheet", "EP01"]) == 0
    assert len(list((d / "cgcut_out").glob("EP01_*.mp4"))) == 6


@pytest.mark.skipif(not HAVE_FFMPEG, reason="ffmpeg 없음")
def test_video_dir_skips_previous_output(sample_dir: Path, tmp_path: Path):
    """--video-dir 안에 cgcut_out이 있어도 결과물을 원본으로 읽지 않음 (bat로 같은 폴더에서 두 번 실행)."""
    import shutil as sh
    d = tmp_path / "proj"
    d.mkdir()
    sh.copy(sample_dir / "CG_list_sample.xlsx", d / "list.xlsx")
    sh.copy(sample_dir / "EP01_sample.mov", d / "EP01_sample.mov")
    args = [str(d / "list.xlsx"), "--video-dir", str(d), "-o", str(d / "cgcut_out"), "--sheet", "EP01", "--codec", "h264"]
    assert cgcut.main(args) == 0
    assert cgcut.main(args) == 0
    assert len(list((d / "cgcut_out").glob("*.mp4"))) == 6


EDL_SAMPLE = """TITLE: EP01 CG
FCM: NON-DROP FRAME

001  A001C003 V     C        05:10:00:00 05:10:03:00 01:00:02:00 01:00:05:00
* FROM CLIP NAME: A001C003_250915.mov
002  BL       V     C        00:00:00:00 00:00:01:00 01:00:05:00 01:00:06:00
003  A002C010 V     C        05:20:00:00 05:20:02:00 01:00:08:08 01:00:10:08
003  A002C011 V     D    024 05:30:00:00 05:30:01:00 01:00:10:08 01:00:11:08
* FROM CLIP NAME: 창문.mov
* TO CLIP NAME: 다음.mov
"""


def test_parse_edl(tmp_path: Path):
    p = tmp_path / "list.edl"
    p.write_text(EDL_SAMPLE, encoding="utf-8")
    rows = cgcut.parse_edl(p)
    assert [(r.no, str(r.tc_in), str(r.tc_out), r.ep) for r in rows] == [
        ("001_A001C003_250915", "01:00:02:00", "01:00:05:00", 1),
        ("003_창문", "01:00:08:08", "01:00:10:08", 1),
    ]
    src = cgcut.parse_edl(p, use_source=True)
    assert str(src[0].tc_in) == "05:10:00:00"
    assert cgcut.load_shots(p)[0].label == "EP01_001_A001C003_250915"


@pytest.mark.skipif(not HAVE_FFMPEG, reason="ffmpeg 없음")
def test_cut_edl(sample_dir: Path, tmp_path: Path):
    import make_sample as ms
    p = tmp_path / "EP01_cg.edl"
    p.write_text(EDL_SAMPLE, encoding="utf-8")
    out = tmp_path / "out"
    assert cgcut.main([str(p), "--video", str(sample_dir / "EP01_sample.mov"), "-o", str(out), "--codec", "h264"]) == 0
    files = sorted(f.name for f in out.glob("*.mp4"))
    assert files == ["EP01_001_A001C003_250915_01-00-02-00.mp4", "EP01_003_창문_01-00-08-08.mp4"]
    # 01:00:02:00~01:00:05:00 (OUT 미포함) = 72프레임, 다음 샷으로 넘어가지 않음
    means = _gray_frames(out / files[0], ms.W, ms.H)
    assert len(means) == 72 and max(means) - min(means) < 3
    assert len(_gray_frames(out / files[1], ms.W, ms.H)) == 48
